import random
import os
from datetime import datetime, timedelta
from docx import Document

# 用户提供的辐射超标地面点位ID列表（384个）
POINT_IDS = [
    55, 56, 57, 58, 85, 86, 87, 88, 115, 116, 117, 118,
    145, 146, 147, 148, 175, 176, 177, 178, 205, 206, 207, 208,
    235, 236, 237, 238, 265, 266, 267, 268, 295, 296, 297, 298,
    325, 326, 327, 328, 355, 356, 357, 358, 385, 386, 387, 388,
    415, 416, 417, 418, 445, 446, 447, 448, 475, 476, 477, 478,
    505, 506, 507, 508, 535, 536, 537, 538, 565, 566, 567, 568,
    595, 596, 597, 598, 625, 626, 627, 628, 655, 656, 657, 658,
    685, 686, 687, 688, 715, 716, 717, 718, 745, 746, 747, 748,
    775, 776, 777, 778, 805, 806, 807, 808, 671, 672, 673, 674,
    675, 676, 677, 678, 679, 680, 681, 682, 683, 684, 701, 702,
    703, 704, 705, 706, 707, 708, 709, 710, 711, 712, 713, 714,
    731, 732, 733, 734, 735, 736, 737, 738, 739, 740, 741, 742,
    743, 744, 761, 762, 763, 764, 765, 766, 767, 768, 769, 770,
    771, 772, 773, 774, 791, 792, 793, 794, 795, 796, 797, 798,
    799, 800, 801, 802, 803, 804, 41, 42, 43, 44, 45, 46,
    47, 48, 49, 50, 71, 72, 73, 74, 75, 76, 77, 78, 79, 80,
    101, 102, 103, 104, 105, 106, 107, 108, 109, 110, 131, 132,
    133, 134, 135, 136, 137, 138, 139, 140, 161, 162, 163, 164,
    165, 166, 167, 168, 169, 170, 191, 192, 193, 194, 195, 196,
    197, 198, 199, 200, 221, 222, 223, 224, 225, 226, 227, 228,
    229, 230, 251, 252, 253, 254, 255, 256, 257, 258, 259, 260,
    281, 282, 283, 284, 285, 286, 287, 288, 289, 290, 311, 312,
    313, 314, 315, 316, 317, 318, 319, 320, 341, 342, 343, 344,
    345, 346, 347, 348, 349, 350, 371, 372, 373, 374, 375, 376,
    377, 378, 379, 380, 401, 402, 403, 404, 405, 406, 407, 408,
    409, 410, 431, 432, 433, 434, 435, 436, 437, 438, 439, 440,
    461, 462, 463, 464, 465, 466, 467, 468, 469, 470, 491, 492,
    493, 494, 495, 496, 497, 498, 499, 500, 521, 522, 523, 524,
    525, 526, 527, 528, 529, 530, 551, 552, 553, 554, 555, 556,
    557, 558, 559, 560, 581, 582, 583, 584, 585, 586, 587, 588,
    589, 590, 611, 612, 613, 614, 615, 616, 617, 618, 619, 620,
    641, 642, 643, 644, 645, 646, 647, 648, 649, 650
]

def generate_random_date(year=2025):
    """生成2025年内的随机日期"""
    start_date = datetime(year, 1, 1)
    end_date = datetime(year, 12, 31)
    days_diff = (end_date - start_date).days
    random_days = random.randint(0, days_diff)
    random_date = start_date + timedelta(days=random_days)
    return random_date

def generate_random_time():
    """生成随机时间（小时、分钟、秒）"""
    hour = random.randint(0, 23)
    minute = random.randint(0, 59)
    second = random.randint(0, 59)
    return hour, minute, second

def format_timedelta(td):
    """将timedelta格式化为'X小时Y分钟Z秒'"""
    total_seconds = int(td.total_seconds())
    hours = total_seconds // 3600
    minutes = (total_seconds % 3600) // 60
    seconds = total_seconds % 60
    return f"{hours}小时{minutes}分钟{seconds}秒"

def find_row_index_by_text(table, text):
    """在表格中查找包含指定文本的行索引"""
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            if text in cell.text:
                return row_idx
    return -1

def generate_report(template_path, output_dir, report_num):
    """生成单个报告文件"""
    # 加载模板文档
    doc = Document(template_path)
    table = doc.tables[0]
    
    # 生成随机日期和时间
    random_date = generate_random_date(2025)
    start_hour, start_minute, start_second = generate_random_time()
    end_hour, end_minute, end_second = generate_random_time()
    
    start_time = datetime(random_date.year, random_date.month, random_date.day, 
                         start_hour, start_minute, start_second)
    end_time = datetime(random_date.year, random_date.month, random_date.day, 
                       end_hour, end_minute, end_second)
    
    if end_time <= start_time:
        hours_to_add = random.randint(1, 4)
        end_time = start_time + timedelta(hours=hours_to_add)
    
    duration = end_time - start_time
    start_time_str = start_time.strftime("%Y-%m-%d %H:%M:%S")
    end_time_str = end_time.strftime("%Y-%m-%d %H:%M:%S")
    duration_str = format_timedelta(duration)
    
    # 随机选择辐射超标地面点位ID（2-20个）
    num_points = random.randint(2, 20)
    selected_ids = random.sample(POINT_IDS, num_points)
    
    # 修改表格内容
    # 修改任务开始时间（行1，列1）
    table.cell(1, 1).text = start_time_str
    # 修改任务结束时间（行2，列1）
    table.cell(2, 1).text = end_time_str
    # 修改任务历时时长（行3，列1）
    table.cell(3, 1).text = duration_str
    # 修改辐射超标地面网格个数（行5，列1）
    table.cell(5, 1).text = str(num_points)
    
    # 查找"辐射超标手套孔名称"所在的行索引
    glove_title_idx = find_row_index_by_text(table, "辐射超标手套孔名称")
    if glove_title_idx == -1:
        glove_title_idx = 12  # 默认值
    
    # 辐射超标地面点位ID行从第9行开始（0-based索引8）
    # 到手套孔标题行之前结束
    start_id_row = 8  # "辐射超标地面点位ID"标题行
    first_data_row = start_id_row + 1  # 第一个数据行
    
    # 清除现有的辐射超标地面点位ID数据行
    # 从第一个数据行开始，到手套孔标题行之前
    for row_idx in range(first_data_row, glove_title_idx):
        if row_idx < len(table.rows):
            # 清空单元格内容
            table.cell(row_idx, 0).text = ""
            table.cell(row_idx, 1).text = ""
    
    # 现在填充新的辐射超标地面点位ID数据
    # 如果需要的行数多于现有行，添加新行
    current_data_rows = glove_title_idx - first_data_row
    if num_points > current_data_rows:
        # 需要添加更多行
        rows_to_add = num_points - current_data_rows
        for _ in range(rows_to_add):
            # 在手套孔标题行之前插入新行
            # 复制最后一行的格式
            new_row = table.add_row()
            # 移动新行到正确位置（在手套孔标题行之前）
            # 由于add_row在末尾添加，我们需要手动移动
            # 这里采用简单方法：先添加，然后通过复制内容来处理
            pass
    
    # 填充数据行
    for i in range(num_points):
        row_idx = first_data_row + i
        if row_idx >= len(table.rows):
            # 如果行不存在，添加新行
            table.add_row()
        
        # 设置点位ID
        table.cell(row_idx, 0).text = str(selected_ids[i])
        # 设置测量值
        table.cell(row_idx, 1).text = "0.12"
    
    # 如果有多余的数据行，清空它们
    total_data_rows = glove_title_idx - first_data_row
    if total_data_rows > num_points:
        for i in range(num_points, total_data_rows):
            row_idx = first_data_row + i
            if row_idx < len(table.rows):
                table.cell(row_idx, 0).text = ""
                table.cell(row_idx, 1).text = ""
    
    # 保存文档
    output_filename = f"普查报告_{report_num:03d}.docx"
    output_path = os.path.join(output_dir, output_filename)
    doc.save(output_path)
    
    return output_path, start_time_str, end_time_str, duration_str, selected_ids

def main():
    # 路径配置
    template_path = r"C:\Users\罗兆东\Desktop\4XTS\AI高价值场景应用\污染普查数据可视化\示例文件\20250701.docx"
    output_dir = r"C:\Users\罗兆东\dist\污染普查数据生成器_100个报告"
    
    # 创建输出目录
    os.makedirs(output_dir, exist_ok=True)
    
    # 生成100个报告
    print("开始生成100个污染普查报告...")
    for i in range(1, 101):
        try:
            output_path, start_time, end_time, duration, selected_ids = generate_report(
                template_path, output_dir, i
            )
            print(f"生成报告 {i:03d}: {os.path.basename(output_path)}")
            print(f"  开始时间: {start_time}")
            print(f"  结束时间: {end_time}")
            print(f"  历时: {duration}")
            print(f"  辐射超标点位ID: {selected_ids}")
            print()
        except Exception as e:
            print(f"生成报告 {i:03d} 时出错: {e}")
            import traceback
            traceback.print_exc()
    
    print(f"报告生成完成！共生成100个文件，保存在: {output_dir}")

if __name__ == "__main__":
    main()